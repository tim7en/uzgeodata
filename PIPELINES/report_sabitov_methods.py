"""Interpret the full local thesis and publish model checks, figures and workbook."""
import hashlib
import json
from collections import defaultdict
from datetime import datetime, timezone

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from build_chirchik_case_studies import ROOT, OUT, read_csv, write_csv, write_json, scores
from analyse_sabitov_methods import prepare, simulate, PARAMETERS

SOURCES = [
    ('Sabitov, T. (2018), full local master thesis', 'storage/Sabitov_Master_ERE_2018.pdf'),
    ('Author-uploaded thesis record', 'https://www.researchgate.net/publication/325033449_HYDROLOGIC_MODELING_OF_GLACIATED_WATERSHED_IN_CENTRAL_ASIA'),
    ('ERA5-Land daily aggregates', 'https://developers.google.com/earth-engine/datasets/catalog/ECMWF_ERA5_LAND_DAILY_AGGR'),
    ('SCS curve-number equations and event-model scope', 'https://www.hec.usace.army.mil/confluence/hmsdocs/hmstrm/canopy-surface-infiltration-and-runoff-volume/infiltration/scs-curve-number-loss-model'),
    ('Hamon equations and calibration', 'https://www.hec.usace.army.mil/confluence/hmsdocs/hmstrm/evaporation-and-transpiration/hamon-method'),
    ('ALOS PALSAR product resolution', 'https://docs.asf.alaska.edu/datasets/palsar/'),
    ('GLIMS snapshot', 'https://developers.google.com/earth-engine/datasets/catalog/GLIMS_20230607'),
    ('Copernicus GLO-30', 'https://developers.google.com/earth-engine/datasets/catalog/COPERNICUS_DEM_GLO30_2024_1'),
]

GAPS = [
    ('Daily forcing and time step', 'pp.15,24–35', 'Implemented', '6,575 actual daily ERA5 values, 2000–2017; daily forcing is not inferred from monthly observations.'),
    ('Four structural models', 'Table 2-1; pp.24–35', 'Adapted', 'Lumped/distributed stores with and without glacier and CN processes; positive ET retained in all structures.'),
    ('Three elevation zones and seasonal lapse rates', 'pp.22,26,34', 'Adapted', 'Exact 2300/3300 m thresholds; open outer bounds retain all terrain; basin ERA5 mean conserved when temperature is redistributed.'),
    ('Independent snow, soil and groundwater storage', 'Eqs.17–22', 'Implemented', 'Daily conservation checks, field-capacity recharge and previous-day groundwater recession; release is 1 minus retention.'),
    ('Five-day antecedent moisture and SCS runoff', 'Eqs.13–16', 'Adapted', 'Prior rain only; dormant/growing thresholds in mm; daily-reset and event-accumulated variants, corrected SI retention.'),
    ('Hamon PET and water-limited AET', 'Eqs.1–2; p.17', 'Adapted', 'Explicit vapour-pressure units; compare corrected Hamon, literal printed equation and ERA5 potential evaporation. No forced 250 mm annual ET.'),
    ('Glacier melt and area', 'Eqs.10–11; pp.28,34', 'Adapted', 'Pskem-specific dated GLIMS inventory, finite ice reservoir, snow shielding; degree-day melt plus annual-cubic sensitivity.'),
    ('Calibration and physical plausibility', 'p.36; Table 2-4', 'Implemented', 'Fixed early calibration/later holdout; daily and monthly scores, seasonal baseline, budget/convergence and boundary parameters disclosed.'),
    ('Flow duration and specific runoff', 'pp.20,38–41', 'Implemented', 'Q50/Q90/Q95, Q90/Q50 proxy and area-normalized discharge; no unsupported recurrence-frequency extrapolation.'),
    ('Monthly flow trends', 'pp.40–42', 'Implemented', 'Kendall/Sen on 2001–2017 monthly flows with Holm correction; exploratory because serial correlation and short record remain.'),
    ('Climate perturbation experiments', 'Chapter 3; pp.58–60', 'Adapted', 'Warming/wetting and warming/drying stress tests with monthly timing and ice-area sensitivity. No date-specific 2050 forecast.'),
    ('Satellite process checks', 'Extension beyond thesis', 'Implemented', 'Held-out MODIS snow and MOD16 ET comparisons retain coverage filters and different spatial supports; not independent field truth.'),
    ('Land-cover confusion matrix', 'p.35; Table 2-3', 'Awaiting reference labels', 'Executable user/producer/overall accuracy evaluation and reference table supplied. No invented labels or transferred 82% accuracy.'),
    ('Soils and spatial curve numbers', 'pp.30–31', 'Awaiting soil evidence', 'CN50/CN90 remain thesis assumptions with sensitivity; Esri classes alone cannot establish hydrologic soil groups.'),
    ('Tributary mass balance and local lapse validation', 'pp.20–22', 'Awaiting observations', 'No simultaneous Oigaing/Maydantal/Charalma discharge or thesis-period paired elevation-confirmed daily meteorology in supplied files.'),
    ('Gauge, ice thickness and recent verification', 'pp.15,28', 'Unresolved', 'Two gauge coordinates audited; fixed 15/30/60 m water-equivalent ice scenarios are not measurements; discharge ends in 2017.'),
]


def extensions(d):
    forcing=read_csv(OUT/'sabitov-daily-forcing.csv');zones=d['inputs']['zones'];factor=d['area_km2']*1000/86400
    dates=[r['date'] for r in forcing];ids=[i for i,t in enumerate(dates) if '2011-01-01'<=t<='2017-12-31']
    scenarios=[]
    for key in ['m3','m4']:
        model=next(m for m in d['models'] if m['key']==key);params=[model['parameters'][p] for p in PARAMETERS]
        for label,dt,pm,gm in [('Historical forcing',0,1,1),('Warmer / wetter: +2.2 C, +5% P',2.2,1.05,1),
            ('Warmer / wetter: +3.1 C, +7% P',3.1,1.07,1),('Warmer / drier: +2 C, -10% P',2,.9,1),
            ('Warmer / drier and half glacier area',2,.9,.5)]:
            perturbed=[{**r,'temperature_c':float(r['temperature_c'])+dt,'precipitation_mm':float(r['precipitation_mm'])*pm} for r in forcing]
            simulation=simulate(prepare(perturbed,zones,key=='m4'),params,glacier_scale=gm)
            monthly=[]
            for month in range(1,13):
                subset=[i for i in ids if int(dates[i][5:7])==month]
                monthly.append({'month':month,'discharge_cms':float(simulation[subset,0].mean()*factor)})
            scenarios.append({'model':key,'scenario':label,'temperature_delta_c':dt,'precipitation_multiplier':pm,'glacier_area_multiplier':gm,
                'monthly':monthly,'mean_discharge_cms':float(simulation[ids,0].mean()*factor),
                'ice_fraction':float(simulation[ids,6].sum()/simulation[ids,0].sum())})
    d['climate_scenarios']=scenarios
    # Paired resampling of whole test years retains within-year dependence.
    daily=read_csv(OUT/'sabitov-daily-predictions.csv')
    train=[r for r in daily if r['phase']=='training' and r['observed_cms']!='']
    climatology={m:float(np.mean([float(r['observed_cms']) for r in train if int(r['date'][5:7])==m])) for m in range(1,13)}
    held=[r for r in daily if r['phase']=='held_out' and r['observed_cms']!='']
    groups=[[r for r in held if r['date'].startswith(str(y))] for y in range(2011,2018)]
    rng=np.random.default_rng(1729);year_samples=rng.integers(0,7,size=(1000,7))
    uncertainty=[]
    for m in d['models']:
        sums=[]
        for group in groups:
            obs=np.array([float(r['observed_cms']) for r in group]);pred=np.array([float(r[m['key']]) for r in group])
            baseline=np.array([climatology[int(r['date'][5:7])] for r in group])
            sums.append([len(obs),float(((pred-obs)**2).sum()),float(((baseline-obs)**2).sum())])
        sums=np.array(sums);sample=sums[year_samples].sum(axis=1)
        delta=np.sqrt(sample[:,1]/sample[:,0])-np.sqrt(sample[:,2]/sample[:,0])
        lo,hi=np.quantile(delta,[.025,.975])
        uncertainty.append({'model':m['key'],'rmse_difference_vs_baseline_cms':m['daily_scores']['rmse']-d['baseline_daily_scores']['rmse'],
            'lower95':float(lo),'upper95':float(hi),'resamples':1000,'years':7})
    d['paired_year_bootstrap']=uncertainty
    d['bootstrap_note']='Paired whole-calendar-year bootstrap of the seven held-out years; 1000 resamples, seed 1729. Negative RMSE difference favours the model. Conditional on the fitted parameters, forcing and boundary; not total predictive uncertainty.'
    exhausted=[m for m in d['models'] if m.get('glacier_exhaustion_date')]
    d['physical_readiness']={'status':'red' if exhausted else 'amber','operational_ready':False,
        'reason':'The assumed glacier reservoirs are exhausted within the historical run. Zero late-period ice contribution is a failed physical assumption, not evidence that Pskem has no glacier runoff.' if exhausted else 'Glacier storage and process fluxes remain unverified.',
        'exhaustion_dates':{m['key']:m['glacier_exhaustion_date'] for m in exhausted}}
    d['climate_scenario_note']='Predeclared stress tests, not probabilities or projections for 2050. +5%/+7% precipitation are rounded analogues of thesis Table 3-1; perturbations also apply during warm-up/calibration years while fitted parameters remain frozen. No dynamic glacier-area retreat model.'
    advanced=json.loads((OUT/'advanced-validation.json').read_text(encoding='utf-8'))
    environment=json.loads((OUT/'environment-modelling.json').read_text(encoding='utf-8'))
    components={r['period']:r for r in d['components']}
    snow=[]
    for row in advanced['snow_monthly']:
        period=row['period']
        if row['sensor']!='combined' or row['elevation_band']!='all' or not '2011-01'<=period<='2017-12' or period not in components:continue
        if row['snow_cover_percent'] is None or row['eligible_days']<10 or row['mean_valid_area_percent']<70:continue
        snow.append({'period':period,'satellite_snow_percent':row['snow_cover_percent'],
            'model_snow_proxy_percent':components[period]['snow_area_proxy_percent'],'eligible_days':row['eligible_days'],
            'mean_valid_area_percent':row['mean_valid_area_percent']})
    et=[]
    for row in environment['surface_monthly']:
        period=row['period']
        if row['product']!='et' or not '2011-01'<=period<='2017-12' or period not in components:continue
        if row['value'] is None or row['covered_days']!=row['expected_days'] or row['mean_valid_area_percent']<70:continue
        et.append({'period':period,'mod16_et_mm':row['value'],'model_et_mm':components[period]['et_mm'],'mean_valid_area_percent':row['mean_valid_area_percent']})
    d['process_checks']={'snow_pairs':snow,'et_pairs':et,
        'snow_scores':scores([r['satellite_snow_percent'] for r in snow],[r['model_snow_proxy_percent'] for r in snow]),
        'et_scores':scores([r['mod16_et_mm'] for r in et],[r['model_et_mm'] for r in et]),
        'note':'Diagnostic comparisons, not measurement-error estimates: monthly mean of a three-zone SWE>1 mm proxy versus MODIS cloud-screened fractional snow; whole-basin model AET versus valid terrestrial MOD16 pixels. Dates and spatial sampling differ. No parameters fitted to these checks.'}
    d['method_gap_matrix']=[dict(zip(['method','thesis_pages','status','implementation'],row)) for row in GAPS]
    d['gauge_audit']['coordinates'][1]['source']='Thesis p.15'
    d['source_hashes'].update({f'PUBLISHED/data/case-studies/{name}':hashlib.sha256((OUT/name).read_bytes()).hexdigest() for name in ['advanced-validation.json','environment-modelling.json','pskem-glims-outlines.geojson']})
    d['code_hashes']={f'PIPELINES/{name}':hashlib.sha256((ROOT/'PIPELINES'/name).read_bytes()).hexdigest() for name in ['extract_sabitov_inputs.py','analyse_sabitov_methods.py','report_sabitov_methods.py']}
    write_json(OUT/'sabitov-methods.json',d)
    for name,rows in [('sabitov-snow-process-check.csv',snow),('sabitov-et-process-check.csv',et),('sabitov-method-gap-matrix.csv',d['method_gap_matrix'])]:
        if rows:write_csv(OUT/name,rows,list(rows[0]))
    flat=[{k:v for k,v in r.items() if k!='monthly'}|m for r in scenarios for m in r['monthly']]
    write_csv(OUT/'sabitov-climate-scenarios.csv',flat,list(flat[0]))


def figures(d):
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,'axes.spines.right':False,'axes.titlelocation':'left'})
    colors=['#2686a3','#54934d','#976fb2','#d58b3e'];paths=[]
    temp=OUT/'sabitov-methods-atlas.tmp.pdf'
    with PdfPages(temp) as pdf:
        def save(fig,name,caption):
            fig.supxlabel(caption,fontsize=8);pdf.savefig(fig);fig.savefig(OUT/f'{name}.png',dpi=170);plt.close(fig);paths.append(f'{name}.png')
        fig,axs=plt.subplots(2,1,figsize=(11.7,8.3),layout='constrained')
        rows=[r for r in d['monthly_predictions'] if '2011-01'<=r['period']<='2017-12']
        x=np.arange(len(rows));axs[0].plot(x,[r['observed'] for r in rows],color='#29394b',label='Observed',lw=2)
        for model,color in zip(d['models'],colors):axs[0].plot(x,[r[model['key']] for r in rows],label=model['key'].upper(),color=color,alpha=.8)
        axs[0].set(xticks=x[::12],xticklabels=[r['period'][:4] for r in rows[::12]],ylabel='Monthly discharge (m³/s)',title='Held-out hydrograph · all four daily model structures');axs[0].legend(ncol=5)
        labels=['Baseline']+[m['key'].upper() for m in d['models']]
        vals=[d['baseline_daily_scores']['nse']]+[m['daily_scores']['nse'] for m in d['models']]
        bars=axs[1].barh(labels,vals,color=['#99a7b0']+colors);axs[1].bar_label(bars,fmt='%.3f',padding=5)
        axs[1].set(xlabel='Daily NSE · larger is better; 1 is a perfect fit',title='Same screened daily observations · 2011–2017')
        save(fig,'sabitov-hindcast','Calibration: 2001–2010; warm-up: 2000. Adaptations with ERA5 weather and a provisional drainage boundary; not thesis score replication.')
        fig,axs=plt.subplots(1,2,figsize=(11.7,8.3),layout='constrained')
        for label,color in [('Observed','#29394b'),('Seasonal baseline','#99a7b0')]+list(zip(['m1','m2','m3','m4'],colors)):
            rows=[r for r in d['flow_duration'] if r['series']==label];axs[0].plot([r['exceedance_percent'] for r in rows],[r['discharge_cms'] for r in rows],label=label,color=color)
        axs[0].set(xlabel='Percentage of days this flow is exceeded',ylabel='Discharge (m³/s)',title='Flow duration · high flows left, low flows right');axs[0].legend()
        rows=d['monthly_trends'];axs[1].errorbar([r['sen_slope_cms_year'] for r in rows],range(12),xerr=[[r['sen_slope_cms_year']-r['slope_lower95'] for r in rows],[r['slope_upper95']-r['sen_slope_cms_year'] for r in rows]],fmt='o',color='#2686a3')
        axs[1].axvline(0,color='#99a7b0');axs[1].set(yticks=range(12),yticklabels=['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec'],xlabel='Sen slope (m³/s per year)',title='Observed monthly flow trends · 2001–2017')
        save(fig,'sabitov-flow-diagnostics','Flow duration uses screened held-out days. Trend intervals are nominal; serial dependence remains, with twelve-test Holm p-values in the workbook.')
        fig,axs=plt.subplots(2,1,figsize=(11.7,8.3),layout='constrained')
        rows=[r for r in d['components'] if '2011-01'<=r['period']<='2017-12'];x=np.arange(len(rows))
        axs[0].stackplot(x,*[[r[k] for r in rows] for k in ['baseflow_mm','quickflow_mm','glacier_melt_mm']],labels=['Groundwater release','Surface runoff','Direct glacier melt'],colors=['#54934d','#2686a3','#976fb2'])
        axs[0].set(ylabel='Monthly runoff depth (mm)',title='M4 modelled runoff pathways');axs[0].legend(ncol=3)
        for key,label,color in [('precipitation_mm','Precipitation','#2686a3'),('et_mm','Actual ET','#d58b3e'),('snow_storage_mm','Snow storage, month end','#976fb2')]:axs[1].plot(x,[r[key] for r in rows],label=label,color=color)
        axs[1].set(ylabel='Water equivalent (mm)',title='M4 water inputs, loss and seasonal snow storage');axs[1].legend(ncol=3)
        for ax in axs:ax.set(xticks=x[::12],xticklabels=[r['period'][:4] for r in rows[::12]])
        save(fig,'sabitov-water-balance','Snowmelt feeds surface/groundwater paths and is not added twice. Glacier contribution is a fixed-inventory model scenario, not observed attribution.')
        fig,ax=plt.subplots(figsize=(11.7,8.3),layout='constrained');rows=d['sensitivity']
        ax.barh([r['scenario'] for r in rows],[r['mean_flow_change_percent'] for r in rows],color='#2686a3');ax.axvline(0,color='#29394b')
        ax.set(xlabel='Change in mean modelled discharge (%)',title='M4 assumptions and parameter sensitivity · frozen calibration')
        save(fig,'sabitov-sensitivity','One-at-a-time perturbations; not causal variable importance or confidence intervals. RMSE changes and conservation checks accompany each scenario.')
        fig,axs=plt.subplots(1,2,figsize=(11.7,8.3),layout='constrained')
        for ax,key in zip(axs,['m3','m4']):
            for row in d['climate_scenarios']:
                if row['model']==key:ax.plot(range(1,13),[r['discharge_cms'] for r in row['monthly']],label=row['scenario'])
            ax.set(xlabel='Month',ylabel='Mean discharge (m³/s)',title=f'{key.upper()} climate stress tests',xticks=[1,3,5,7,9,11]);ax.legend(fontsize=7)
        save(fig,'sabitov-climate-stress','Historical daily sequence perturbed with frozen parameters. These are stress tests, not calibrated 2050 projections or scenario probabilities.')
        fig,axs=plt.subplots(1,2,figsize=(11.7,8.3),layout='constrained')
        for ax,rows,xkey,ykey,label,unit in [(axs[0],d['process_checks']['snow_pairs'],'satellite_snow_percent','model_snow_proxy_percent','Snow','%'),(axs[1],d['process_checks']['et_pairs'],'mod16_et_mm','model_et_mm','Evapotranspiration','mm/month')]:
            if rows:
                ax.scatter([r[xkey] for r in rows],[r[ykey] for r in rows],s=15,alpha=.6,color='#2686a3');hi=max(max(r[xkey],r[ykey]) for r in rows)
                ax.plot([0,hi],[0,hi],ls='--',color='#99a7b0')
            ax.set(xlabel=f'Satellite estimate ({unit})',ylabel=f'M4 model ({unit})',title=f'{label} process check · {len(rows)} pairs')
        save(fig,'sabitov-process-checks','2011–2017; coverage thresholds retained. Three-zone snow proxy and whole-basin AET have different spatial support from satellite valid pixels.')
    temp.replace(OUT/'sabitov-methods-atlas.pdf')
    return paths


def review(d):
    ice=sum(z['glacier_area_km2'] for z in d['inputs']['zones'])
    lines=['# Sabitov (2018): full-thesis methodology audit and implemented Pskem experiments','',
        'This review uses the supplied 76-page file `storage/Sabitov_Master_ERE_2018.pdf`, read in full. Page references below are printed thesis pages; add nine for the PDF page number. Timur Sabitov is the sole thesis author; supervisors are not listed here as coauthors. The earlier abstract-only review is superseded.','',
        '## Finding','',
        'Our earlier case study had monthly water-balance, random-forest and Bayesian comparisons but lacked the daily hydrologic structure central to the thesis. The extension now executes four daily model adaptations, elevation-dependent snow and glacier processes, five-day antecedent rainfall, SCS runoff, distinct soil and groundwater stores, Hamon ET, flow-duration analysis, monthly trend diagnostics and climate stress tests. It retains the existing statistical models as separate monthly experiments.','',
        'A more elaborate model is not automatically more credible. Daily and monthly scores answer different questions; a good discharge curve can conceal wrong snow storage, ET or glacier melt. The new outputs expose these checks and limitations rather than declaring the best score operational.','',
        '## What the thesis actually did','',
        'Chapter 2 used local Pskem climate and streamflow, glacier information, DEM topography and a six-class unsupervised Landsat classification. Model 1 was a lumped snow/storage balance; Model 2 added three elevation zones and empirical glacier melt; Model 3 added SCS runoff and separate unsaturated/saturated stores in a lumped basin; Model 4 combined the SCS structure with elevation zones. The thesis assumed CN near 50 downstream and 90 in the upper rocky/glaciated zone, seasonal lapse-rate magnitudes of 6–6.7 °C/km, and an empirical glacier temperature at 3643 m. These were partly calibrated/assumed properties, not distributed observations (pp.24–35).','',
        'The last two observed years were used for calibration after a first-year warm-up (p.36). There is no independent later test period in that description. Table 2-4 (p.53) reports daily NSE 0.70, 0.23, 0.76 and 0.29 for Models 1–4; monthly NSE is 0.90, 0.24, 0.77 and 0.69. The abstract gives monthly R² values 0.84, 0.77, 0.93 and 0.85; body passages use correlation terminology. Do not equate r, R² and NSE. Model 1 achieves its attractive fit with zero ET; that is not adopted as a physically acceptable calibration strategy here.','',
        'The simulation dates also need author/source-code clarification: the abstract describes water years 2013–2015, a results passage says October 2013–September 2015, and Table 3-1 gives 1095 daily values. Without original forcing and code, an exact numerical reproduction cannot be asserted.','',
        '## Method-by-method gap matrix','',
        '| Method | Thesis location | Current status | Implementation / remaining requirement |','| --- | --- | --- | --- |']
    lines += [f'| {a} | {b} | {c} | {e} |' for a,b,c,e in GAPS]
    lines += ['', '## Dimensional corrections and deliberate departures','',
        '1. **SCS depth and threshold.** Eq.16 prints `2540/CN − 25.4`, which is a centimetre expression. The implementation consistently uses `S = 25400/CN − 254` in millimetres; direct runoff is zero when liquid input is at or below `0.2S`. Depth is converted to m³/s only after area weighting. Five-day thresholds are 13/28 mm dormant and 36/53 mm growing; April–September is our declared season assumption. Antecedent rain excludes today. The basic curve-number relationship is event based. Daily resetting is retained as a thesis-style approximation and compared with an event-accumulated variant that resets after a day with at most 0.1 mm liquid input. The extension to snowmelt is empirical. [USACE SCS method]('+SOURCES[3][1]+').','',
        '2. **Groundwater and ET accounting.** Eq.20 releases `(1 − Kb) × SAT`, so Kb=0.99 releases 1% of previous groundwater storage per day. ET is removed once from available soil water, before field-capacity recharge; it is not subtracted from discharge again. Each zone checks precipitation minus ET minus runoff against changes in snow, soil, groundwater and glacier-ice stores.','',
        '3. **Hamon units.** Thesis Eqs.1–2 mix a cm/day PET label with a saturation-vapour-pressure formula whose output is kPa. Our primary form uses `es = 0.6108 exp(17.27T/(237.3+T))` kPa, vapour density `2167 es/(T+273.15)` g/m³ and `PET = 0.1651 (daylight/12) density` mm/day. The factor 2167 explicitly converts kPa to hPa relative to the usual 216.7 expression. Daylight is astronomical, and T≤0 gives zero PET as in the thesis. Literal printed Hamon (including cm-to-mm conversion) and ERA5 potential evaporation are separate sensitivities. Crop coefficient remains positive and ET cannot exceed available water. The thesis 250 mm/year literature target is not treated as an observation. [USACE Hamon method]('+SOURCES[4][1]+').','',
        '4. **Glacier melt.** Eq.10 labels the cubic ablation expression as m³/s, whereas Eq.11 requires annual ablation in mm/year. Applying the annual relationship directly as daily flow is dimensionally unsafe. The main adaptation instead uses 6 mm/°C/day ice melt at the stated 3643 m reference altitude, capped by remaining ice and shielded by remaining seasonal snow. The annual-cubic relationship divided by 365.25 is a sensitivity only; using daily temperature in it remains an empirical approximation. Snow melt uses a declared 4.5 mm/°C/day assumption, interpreting the thesis 0.45 coefficient as cm/°C/day; 3 and 6 mm alternatives are tested. No claim is made that these melt factors were locally measured.','',
        '5. **Initial ice and geometry.** Glacier outlines are fixed, and initial water-equivalent ice depth is assumed to be 30 m. Alternatives of 15/60 m and half mapped glacier area expose this uncertainty. Those depths are scenario parameters, not measured thickness or geodetic mass balance. Direct glacier drainage is a thesis assumption; separate routing/tracer observations are absent.','',
        '6. **Elevation and land cover.** Open lower/upper elevation bounds avoid dropping terrain outside the thesis 1251–4300 m domain. Temperature redistribution preserves the ERA5 basin mean around the DEM mean; it does not correct model orography or substitute unverified station elevations. Pskem station height is variously 1251/1254 m and the thesis elevation difference is arithmetically inconsistent; local lapse validation awaits resolved metadata. No unsupported positive precipitation gradient is imposed, especially because thesis higher Oigaing is drier. The 10 m Esri class series cannot distinguish every thesis mixed class or independently assign soil groups.','',
        '7. **DEM resolution and metrics.** ASF PALSAR RTC 12.5 m pixel spacing does not establish a native 12.5 m elevation model; source DEM resolution must be checked. The current 30 m Copernicus product is retained, with no artificial precision from resampling. Standard NSE uses paired simulated minus observed values; the apparent mean-observation substitution in thesis Eq.24 is not copied. Bias here is simulation minus observation. [ASF PALSAR documentation]('+SOURCES[5][1]+').','',
        '## Actual inputs and spatial audit','',
        f"Earth Engine supplied {d['inputs']['rows']:,} daily basin records from {d['inputs']['start']} through {d['inputs']['end']}. Three 30 m raster zones total {d['area_km2']:.2f} km², about 0.13% above the sum of nominal HydroBASINS areas (2626.9 km²); raster area is used consistently for the new depth-to-flow conversion. The thesis reports 2540 km². These boundaries are not interchangeable. [ERA5 daily catalogue]({SOURCES[2][1]}).",'',
        f"The old headwater glacier export had no Pskem coverage. A catchment-specific GLIMS query selected {d['inputs']['glacier_features']} latest-per-ID outlines, removed matching internal rock and unioned overlaps. Rasterized ice totals {ice:.2f} km² ({100*ice/d['area_km2']:.2f}% of the candidate basin). Outline dates are {', '.join(d['inputs']['glacier_outline_dates'])}; the 2023 catalogue date is not a 2023 survey. This is a static hindcast inventory, including outlines surveyed after the earliest forcing years. [GLIMS catalogue]({SOURCES[6][1]}).",'',
        '| Elevation zone | Area km² | Mean elevation m | Glacier area km² |','| --- | ---: | ---: | ---: |']
    lines += [f"| {z['label']} | {z['area_km2']:.2f} | {z['mean_elevation_m']:.1f} | {z['glacier_area_km2']:.2f} |" for z in d['inputs']['zones']]
    lines += ['',f"The thesis gauge position (p.15), 41°47′N 70°13′E, is {d['gauge_audit']['coordinate_separation_m']/1000:.2f} km from the existing coordinate. Both were checked against held HydroRIVERS geometry:",'', '| Coordinate source | Longitude | Latitude | Closest reach | Offset m | Upstream km² |','| --- | ---: | ---: | --- | ---: | ---: |']
    for p in d['gauge_audit']['coordinates']:
        n=p['nearest_reaches'][0];lines.append(f"| {p['source']} | {p['longitude']:.6f} | {p['latitude']:.6f} | {n['reach_id']} | {n['distance_m']:.1f} | {n['upstream_area_km2']:.1f} |")
    lines += ['', 'Neither coordinate was silently substituted. The complete nearest-reach alternatives are in `sabitov-methods.json`; original station metadata, river identity and a reviewed outlet delineation remain necessary.','',
        '## Held-out results','',f"Daily calibration uses {d['split']['training_days']:,} screened observations in 2001–2010; evaluation uses {d['split']['held_out_days']:,} in 2011–2017. The invalid 2015-02-29 row and three manifest-flagged values are quarantined. Monthly comparisons use the same observed dates on both sides and at least 90% daily coverage. Missing gauge days are never filled.",'',
        '| Model | Daily RMSE m³/s | Daily NSE | Monthly NSE | Modelled ice share | Mean AET mm/year |','| --- | ---: | ---: | ---: | ---: | ---: |',
        f"| Training seasonal baseline | {d['baseline_daily_scores']['rmse']:.2f} | {d['baseline_daily_scores']['nse']:.3f} | — | — | — |"]
    for m in d['models']:lines.append(f"| {m['label']} | {m['daily_scores']['rmse']:.2f} | {m['daily_scores']['nse']:.3f} | {m['monthly_scores']['nse']:.3f} | {100*m['held_out_ice_fraction_of_modelled_flow']:.2f}% | {m['held_out_et_mm_year']:.1f} |")
    lines += ['', '**Physical failure exposed:** '+d['physical_readiness']['reason']+' Reservoir exhaustion dates: '+', '.join(f'{k.upper()}: {v}' for k,v in d['physical_readiness']['exhaustion_dates'].items())+'. The 60 m scenario retains more ice and changes flow; selecting a thickness to improve held-out scores would not validate it. No glacier contribution or future water-supply conclusion is accepted from the depleted runs.','',
        'These are new adaptations evaluated on a different period and forcing, not recovered thesis results. Snow, ET, geometry and calibration constraints can change the ordering. The main figures show every predefined model; the holdout is not used to tune a winning structure. Optimizer budget/termination and parameter-bound flags are exported. A bounded local search does not establish a global optimum.','',
        f"Maximum model water-balance residual is {max(m['max_abs_mass_balance_residual_mm'] for m in d['models']):.2g} mm/day. This confirms numerical bookkeeping, not that flux magnitudes are correct. Fixed-parameter sensitivity experiments separate assumptions about CN, rainfall, ET, melt factors and glacier storage from fitted model performance.",'',
        '## Process checks, low flows and climate scenarios','',
        f"The M4 snow-area proxy has {d['process_checks']['snow_scores']['n']} eligible held-out MODIS monthly pairs; its RMSE is {d['process_checks']['snow_scores']['rmse']:.2f} percentage points. The ET check has {d['process_checks']['et_scores']['n']} eligible MOD16 pairs and RMSE {d['process_checks']['et_scores']['rmse']:.2f} mm/month. These differences include spatial/temporal support effects: three binary elevation zones versus fractional clear-sky snow, and whole-basin AET versus valid vegetated/terrestrial pixels. They reveal process mismatch, not independent sensor accuracy. Neither check influenced calibration.",'',
        f"Held-out observed Q50/Q90/Q95 are {d['low_flow']['q50_cms']:.2f}, {d['low_flow']['q90_cms']:.2f}, {d['low_flow']['q95_cms']:.2f} m³/s. Q90/Q50 is {d['low_flow']['q90_q50_proxy']:.3f}; it is a low-flow persistence proxy, not hydrograph-separated baseflow. Specific runoff is {d['low_flow']['specific_discharge_l_s_km2']:.2f} L/s/km² under the provisional area. Annual recurrence estimates and tributary closure are not manufactured from missing data or copied from the thesis table with inconsistent ratios.",'',
        'Monthly Kendall/Sen analyses use the available 2001–2017 record and twelve-test Holm correction, not the thesis 1965–2015 record. Autocorrelation and the short record limit inference; nominal intervals and p-values are descriptive screening. No trend is labelled anthropogenic attribution.','',
        'Chapter 3 explored warming and altered precipitation. We implement rounded analogues (+2.2 °C/+5% P and +3.1 °C/+7% P), a +2 °C/−10% P case, and the latter with half glacier area. Both M3 and M4 retain frozen calibration and the same historical day sequence. Monthly hydrographs expose timing changes and structural disagreement. These are conditional stress tests, not calibrated probabilities, CMIP6 ensembles or a forecast for 2050.','',
        '## Remaining evidence required','',
        'The study can now execute the missing process methods, but additional measurements are still needed for defensible physical attribution: verified gauge position and drainage area; recent discharge/rating curves; daily station precipitation and temperature at confirmed elevations; glacier thickness or geodetic balance and area histories; hydrologic soil groups; independent land-cover labels; and tributary flow or routing constraints.','',
        'The reference CSV `CASE_STUDIES/landcover-reference-labels.csv` deliberately starts with no labelled samples. Add independently interpreted, dated samples with unique IDs and mapped/reference classes; rerunning computes the confusion matrix and user/producer/sample-overall accuracy. A stratified sample requires sampling probabilities and area-adjusted estimation before claiming basin-wide accuracy. The thesis 110 polygons and 82% overall accuracy do not transfer to Esri or to this basin/time automatically.','',
        'Dynamic glacier geometry, full surface-energy-balance melt, spatially supported soil/CN estimates, event routing and ensemble uncertainty remain separate extensions. Existing Bayesian/forest experiments are retained; they do not validate these physical parameters. Scientific readiness remains blocked even when every processing stage finishes.','',
        '## Reproduce and inspect','',
        'Run `npm run cases:sabitov:inputs` once for authenticated Earth Engine downloads, then `npm run cases:sabitov` for offline models, checks and artifacts. `npm run cases:update:offline` also includes these stages. Input/code SHA-256 hashes, source image IDs, projection, masks and dates are exported with `sabitov-methods.json` and `sabitov-artifacts.manifest.json`.', '',
        'Outputs: `sabitov-methods-atlas.pdf`, six figure PNGs, `sabitov-methods-analysis.xlsx`, daily forcing/prediction tables, process components, flow-duration curves, month trends, scenario curves and the method gap matrix. The website includes these results under “Sabitov methodology”.','', '## Sources','']
    lines += [f'- [{title}]({url})' if url.startswith('http') else f'- {title}: `{url}`; SHA-256 `{d["thesis"]["sha256"]}`.' for title,url in SOURCES]
    (OUT/'sabitov-2018-review.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')


def workbook(d):
    wb=Workbook();ws=wb.active;ws.title='Read me';ws.append(['Property','Interpretation'])
    for row in [('Status',d['status']),('Calibration','2000 warm-up; 2001–2010 fit; 2011–2017 holdout'),('Thesis',d['thesis']['local_file']),('Units','Column names retain units. Missing is blank, not zero.'),('Ice','Fixed historical outlines; 30 m assumed water-equivalent initial depth'),('Scenarios',d['climate_scenario_note'])]:ws.append(row)
    names=['sabitov-daily-forcing.csv','sabitov-daily-predictions.csv','sabitov-model-components.csv','sabitov-flow-duration.csv','sabitov-monthly-trends.csv','sabitov-snow-process-check.csv','sabitov-et-process-check.csv','sabitov-climate-scenarios.csv','sabitov-method-gap-matrix.csv']
    for name in names:
        ws=wb.create_sheet(name.replace('sabitov-','').replace('.csv','')[:31]);rows=read_csv(OUT/name)
        if not rows:continue
        ws.append(list(rows[0]))
        for r in rows:
            values=[]
            for k,v in r.items():
                if v=='':values.append(None);continue
                try:values.append(float(v))
                except ValueError:values.append(v)
            ws.append(values)
    ws=wb.create_sheet('Scores and parameters');ws.append(['Model','Daily NSE','Daily RMSE','Monthly NSE',*PARAMETERS,'Optimizer converged','Boundary parameters'])
    for m in d['models']:ws.append([m['key'],m['daily_scores']['nse'],m['daily_scores']['rmse'],m['monthly_scores']['nse'],*[m['parameters'][k] for k in PARAMETERS],m['optimizer']['success'],', '.join(m['parameters_at_bounds'])])
    ws=wb.create_sheet('Sensitivity');ws.append(['Scenario','RMSE change m3/s','Mean flow change %','Balance residual mm'])
    for r in d['sensitivity']:ws.append([r['scenario'],r['rmse_change_cms'],r['mean_flow_change_percent'],r['max_abs_mass_balance_residual_mm']])
    ws=wb.create_sheet('Paired year bootstrap');ws.append(['Model','RMSE difference vs baseline','Lower 95%','Upper 95%','Resamples','Years'])
    for r in d['paired_year_bootstrap']:ws.append(list(r.values()))
    ws=wb.create_sheet('Sources');ws.append(['Title','Location'])
    for row in SOURCES:ws.append(row)
    for ws in wb:
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17495B')
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(65,max(18,len(str(col[0].value))+2))
    temp=OUT/'sabitov-methods-analysis.tmp.xlsx';wb.save(temp);temp.replace(OUT/'sabitov-methods-analysis.xlsx')


def main():
    d=json.loads((OUT/'sabitov-methods.json').read_text(encoding='utf-8'))
    extensions(d);paths=figures(d);review(d);workbook(d)
    names=paths+['sabitov-methods.json','sabitov-2018-review.md','sabitov-methods-atlas.pdf','sabitov-methods-analysis.xlsx']
    write_json(OUT/'sabitov-artifacts.manifest.json',{'generated_at':datetime.now(timezone.utc).isoformat(),
        'input_hashes':d['source_hashes'],'code_hashes':d['code_hashes'],'outputs':{name:hashlib.sha256((OUT/name).read_bytes()).hexdigest() for name in names}})
    print('Published full-thesis review, six-page figure atlas, workbook, scenarios and satellite process checks.',flush=True)


if __name__=='__main__':main()
