"""Turn the regional air-temperature/precipitation and soil-temperature
workbooks into one long-format monthly table, matched to the meteorological
station network.

`storage/air_temp_precip/*.xlsx` holds seven regional deliveries, one sheet
per station named "<WMO code> <station name>" with three fixed blocks - Air
temperature, Precipitation, Pedya index - each headed by its own "Years" row
ordered October through September. The printed year is interpreted as the
ending year, cross-checked against the separate Tashkent/Pskem calendar records;
this convention remains inferred for other stations. Negative precipitation quarantines a sheet as
suspect block labels; it does not justify automatically exchanging variables.
Raw values remain in CSV; quarantined values are null in the browser JSON.

`storage/soil_temp/Температура почвы.xlsx` holds thirteen sheets, one per
oblast, each carrying several station blocks in the calendar-year
convention (January-December). Its own "ND" marker means not measured.

Station names arrive in Cyrillic or Latin depending on the file; both are
matched to `hydromet-meteo-network.csv` (built by
build_hydromet_station_network.py) after the same homoglyph repair and
transliteration used there, so a station is one row whichever alphabet its
delivery used.

    python PIPELINES/build_hydromet_station_network.py   # first, if not already built
    python PIPELINES/build_regional_climate_observations.py
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import math
import hashlib
from datetime import datetime, timezone
from pathlib import Path

from lib_cyrillic import clean_text, normalize_cyrillic, transliterate

ROOT = Path(__file__).resolve().parent.parent
AIR_TEMP_DIR = ROOT / "storage/air_temp_precip"
SOIL_TEMP_PATH = ROOT / "storage/soil_temp/Температура почвы.xlsx"
PUBLISHED_DIR = ROOT / "PUBLISHED/data/hydroclimate"
METEO_NETWORK = PUBLISHED_DIR / "hydromet-meteo-network.csv"

OBSERVATIONS_CSV = PUBLISHED_DIR / "regional-climate-monthly.csv"
OBSERVATIONS_JSON = PUBLISHED_DIR / "regional-climate-monthly.json"
MANIFEST = PUBLISHED_DIR / "regional-climate-monthly.manifest.json"

# Block label -> (variable name, unit). "Pedya index" is read like the others
# but its columns are empty in every sheet in this delivery; it is kept so a
# future delivery that fills it in is picked up without a code change.
BLOCKS = {
    "Air temperature": ("air_temperature_mean", "degC"),
    "Precipitation": ("precipitation_total", "mm"),
    "Pedya index": ("ped_drought_index", "index"),
}
# Printed year is the ending year (October of previous year to September).
# Evidence: calendar_alignment_audit against independently imported observations.
WATER_YEAR_MONTHS = {
    "October": (10, -1), "November": (11, -1), "December": (12, -1),
    "January": (1, 0), "February": (2, 0), "March": (3, 0), "April": (4, 0),
    "May": (5, 0), "June": (6, 0), "July": (7, 0), "August": (8, 0), "September": (9, 0),
}
DATE_CHECKED_STATIONS = {'uz:station/meteo-413694', 'uz:station/meteo-419704'}
CALENDAR_MONTHS_RU = {
    "Январь": 1, "Февраль": 2, "Март": 3, "Апрель": 4, "Май": 5, "Июнь": 6,
    "Июль": 7, "Август": 8, "Сентябрь": 9, "Октябрь": 10, "Ноябрь": 11, "Декабрь": 12,
}
NO_DATA = {"nd", "н/д", "-", "—", ""}

# Physically-grounded plausibility bounds. Precipitation must be non-negative;
# monthly means of temperature this far outside Uzbekistan's record are not
# corrected, only flagged, per this project's convention of never overwriting
# a source's own fact.
RANGES = {
    "air_temperature_mean": (-35.0, 50.0),
    "precipitation_total": (0.0, 500.0),
    "soil_temperature_mean": (-30.0, 55.0),
    "ped_drought_index": (-25.0, 25.0),
}

# Oblast sheet name (as delivered, with its known typo) -> the region name
# already canonical in this project (GEODATA/uzb_admbnda_adm1_2018b ADM1_EN).
OBLAST_EN = {
    "РК": "Republic of Karakalpakstan",
    "Хорезм": "Khorezm",
    "Навоийская": "Navoi",
    "Бухарская": "Bukhara",
    "Самаркандская": "Samarkand",
    "Кашкадарьинская": "Kashkadarya",
    "Сурхандарьинская": "Surkhandarya",
    "Джизаксакая": "Dzhizak",  # sheet name misspells Джизакская
    "Сырдарьинская": "Syrdarya",
    "Ташкентская": "Tashkent",
    "Андижанская": "Andizhan",
    "Наманганская": "Namangan",
    "Ферганская": "Fergana",
}


from hydromet.io import write_csv, write_json


def number(value):
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value) if math.isfinite(float(value)) else None
    text = str(value).strip().replace(",", ".")
    if text.lower() in NO_DATA:
        return None
    try:
        parsed = float(text)
        return parsed if math.isfinite(parsed) else None
    except ValueError:
        return None


def simple_key(text: str) -> str:
    """A name reduced to bare lowercase Latin letters, for cross-alphabet matching."""
    latin = transliterate(normalize_cyrillic(text))
    return re.sub(r"[^a-z0-9]", "", latin.lower())


def load_station_index() -> dict[str, dict]:
    """Meteo network rows keyed by their simplified Latin name."""
    index: dict[str, dict] = {}
    if not METEO_NETWORK.exists():
        return index
    with METEO_NETWORK.open(encoding="utf-8", newline="") as handle:
        for row in csv.DictReader(handle):
            key = simple_key(row["name_latin"])
            if key:
                if key in index:
                    raise ValueError(f"Ambiguous station name: {key}")
                index[key] = row
    return index


def match_station(name_latin: str, station_index: dict[str, dict]):
    """Match exact normalized names only; fuzzy suggestions require review."""
    key = simple_key(name_latin)
    exact = station_index.get(key)
    if exact is not None:
        return exact, "exact"
    import difflib
    candidates = difflib.get_close_matches(key, station_index.keys(), n=1, cutoff=0.82)
    if candidates:
        return None, "fuzzy_candidate_requires_review"
    return None, "none"


def read_air_temp_precip(path: Path, station_index: dict, retrieved: str,
                          rows: list[dict], swaps: list[str], unmatched: set[str]):
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.sheetnames:
        grid = list(workbook[sheet].iter_rows(values_only=True))
        if len(grid) < 3:
            continue
        header, months_row = grid[0], grid[1]

        match = re.match(r"\s*(\d+)\s+(.+)$", sheet)
        wmo_code = match.group(1) if match else ""
        name_latin_raw = match.group(2).strip() if match else sheet.strip()

        # Locate each block by its own title, not by an assumed column offset:
        # this survives a sheet that drops the (always-empty) Pedya index block.
        blocks = []  # (variable, unit, years_col, month_columns: {colname: (month, yearoffset)})
        for col, title in enumerate(header):
            title = clean_text(title)
            if title not in BLOCKS:
                continue
            variable, unit = BLOCKS[title]
            month_columns = {}
            for offset in range(1, 13):
                month_name = clean_text(months_row[col + offset]) if col + offset < len(months_row) else ""
                mapping = WATER_YEAR_MONTHS.get(month_name)
                if mapping:
                    month_columns[col + offset] = mapping
            blocks.append({"variable": variable, "unit": unit, "years_col": col, "months": month_columns})

        # Negative precipitation signals suspect labels, not proof of a swap.
        precip_block = next((b for b in blocks if b["variable"] == "precipitation_total"), None)
        temp_block = next((b for b in blocks if b["variable"] == "air_temperature_mean"), None)
        swapped = False
        if precip_block and temp_block:
            for data_row in grid[2:]:
                for col in precip_block["months"]:
                    value = number(data_row[col]) if col < len(data_row) else None
                    if value is not None and value < 0:
                        swapped = True
                        break
                if swapped:
                    break
            if swapped:
                swaps.append(f"{path.name} / {sheet}")

        station, match_method = match_station(name_latin_raw, station_index)
        if station is None:
            unmatched.add(f"{path.name}:{sheet}")

        for block in blocks:
            for row_number, data_row in enumerate(grid[2:], start=3):
                years_cell = data_row[block["years_col"]] if block["years_col"] < len(data_row) else None
                water_year = number(years_cell)
                if water_year is None or not (1900 < water_year < 2100):
                    continue
                water_year = int(water_year)
                for col, (month, year_offset) in block["months"].items():
                    value = number(data_row[col]) if col < len(data_row) else None
                    if value is None:
                        continue
                    calendar_year = water_year + year_offset
                    flag = quality_flag(block["variable"], value)
                    if swapped:
                        flag = "suspect_block_labels"
                    rows.append({
                        "station_wmo_code": wmo_code,
                        "station_name_raw": name_latin_raw,
                        "station_name_latin": transliterate(normalize_cyrillic(name_latin_raw)),
                        "entity_id": station["entity_id"] if station else "",
                        "match_method": match_method,
                        "region_en": "",
                        "water_year": water_year,
                        "calendar_year": calendar_year,
                        "month": month,
                        "variable": block["variable"],
                        "value": f"{value:g}",
                        "unit": block["unit"],
                        "quality_flag": flag,
                        "block_swap_corrected": "no",
                        "source_file": path.name,
                        "source_sheet": sheet,
                        "source_cell": f"{openpyxl.utils.get_column_letter(col + 1)}{row_number}",
                        "date_status": "cross_checked_ending_year" if station and station['entity_id'] in DATE_CHECKED_STATIONS else "inferred_ending_year",
                        "retrieved_at": retrieved,
                    })
    workbook.close()


def quality_flag(variable: str, value: float) -> str:
    bounds = RANGES.get(variable)
    if bounds and not (bounds[0] <= value <= bounds[1]):
        return "suspect_out_of_range"
    return "ok"


def read_soil_temperature(path: Path, station_index: dict, retrieved: str,
                           rows: list[dict], unmatched: set[str]):
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook.sheetnames:
        grid = list(workbook[sheet].iter_rows(values_only=True))
        if len(grid) < 3:
            continue
        header, months_row = grid[0], grid[1]
        oblast_raw = clean_text(sheet)
        region_en = OBLAST_EN.get(oblast_raw, transliterate(normalize_cyrillic(oblast_raw)))

        blocks = []
        block_width = 14
        for col in range(0, len(header), block_width):
            name_raw = clean_text(header[col]) if col < len(header) else ""
            if not name_raw:
                continue
            month_columns = {}
            for offset in range(1, 13):
                month_name = clean_text(months_row[col + offset]) if col + offset < len(months_row) else ""
                month = CALENDAR_MONTHS_RU.get(month_name)
                if month:
                    month_columns[col + offset] = month
            blocks.append({"years_col": col, "name_raw": name_raw, "months": month_columns})

        for block in blocks:
            name_clean = normalize_cyrillic(block["name_raw"])
            name_latin = transliterate(name_clean)
            station, match_method = match_station(name_latin, station_index)
            if station is None:
                unmatched.add(f"{path.name}:{sheet}:{block['name_raw']}")
            for row_number, data_row in enumerate(grid[2:], start=3):
                year = number(data_row[block["years_col"]]) if block["years_col"] < len(data_row) else None
                if year is None or not (1900 < year < 2100):
                    continue
                year = int(year)
                for col, month in block["months"].items():
                    value = number(data_row[col]) if col < len(data_row) else None
                    if value is None:
                        continue
                    flag = quality_flag("soil_temperature_mean", value)
                    rows.append({
                        "station_wmo_code": "",
                        "station_name_raw": block["name_raw"],
                        "station_name_latin": name_latin,
                        "entity_id": station["entity_id"] if station else "",
                        "match_method": match_method,
                        "region_en": region_en,
                        "water_year": "",
                        "calendar_year": year,
                        "month": month,
                        "variable": "soil_temperature_mean",
                        "value": f"{value:g}",
                        "unit": "degC",
                        "quality_flag": flag,
                        "block_swap_corrected": "no",
                        "source_file": path.name,
                        "source_sheet": sheet,
                        "source_cell": f"{openpyxl.utils.get_column_letter(col + 1)}{row_number}",
                        "date_status": "explicit_calendar_year",
                        "retrieved_at": retrieved,
                    })
    workbook.close()


def calendar_alignment_audit(rows):
    """Report alternative year alignments; never silently optimise dates."""
    reference = PUBLISHED_DIR / 'pskem-station-monthly.csv'
    if not reference.exists():
        return {'status': 'reference_missing', 'comparisons': []}
    with reference.open(encoding='utf-8', newline='') as handle:
        lookup = {(r['station_id'], r['variable'], int(r['year']), int(r['month'])): float(r['value']) for r in csv.DictReader(handle)}
    from collections import defaultdict
    groups = defaultdict(list)
    for row in rows:
        if row['variable'] not in ['air_temperature_mean', 'precipitation_total'] or row['quality_flag'] != 'ok':
            continue
        for shift in [-1, 0, 1]:
            key = (row['entity_id'], row['variable'], row['calendar_year']+shift, row['month'])
            if key in lookup:
                group = (row['entity_id'], row['variable'], 'Jan-Sep' if row['month'] <= 9 else 'Oct-Dec', shift)
                groups[group].append(abs(float(row['value'])-lookup[key]))
    return {'reference': reference.name, 'reference_sha256': hashlib.sha256(reference.read_bytes()).hexdigest(),
            'selected': 'ending_year', 'note': 'Zero shift is the selected ending-year mapping. +1 reproduces the former starting-year mapping. Differences in observations still require review.',
            'comparisons': [{'station_id': k[0], 'variable': k[1], 'season': k[2], 'year_shift': k[3],
                             'n': len(v), 'mae': sum(v)/len(v)} for k,v in sorted(groups.items())]}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.parse_args()

    retrieved = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    station_index = load_station_index()
    if not station_index:
        print("Warning: hydromet-meteo-network.csv is missing; run build_hydromet_station_network.py first. "
              "Continuing without station matching.")

    rows: list[dict] = []
    swaps: list[str] = []
    unmatched: set[str] = set()

    for path in sorted(AIR_TEMP_DIR.glob("*.xlsx")):
        before = len(rows)
        read_air_temp_precip(path, station_index, retrieved, rows, swaps, unmatched)
        print(f"{path.name}: {len(rows) - before:,} monthly values")

    if SOIL_TEMP_PATH.exists():
        before = len(rows)
        read_soil_temperature(SOIL_TEMP_PATH, station_index, retrieved, rows, unmatched)
        print(f"{SOIL_TEMP_PATH.name}: {len(rows) - before:,} monthly values")

    rows.sort(key=lambda row: (row["source_file"], row["station_name_latin"], row["variable"],
                                row["calendar_year"], row["month"]))
    # Retain every source record; duplicate conflicts cannot silently overwrite.
    from collections import defaultdict
    groups = defaultdict(list)
    for row in rows:
        identity = row['entity_id'] or (row['source_file'], row['source_sheet'], row['station_name_latin'])
        groups[(str(identity), row['variable'], row['calendar_year'], row['month'])].append(row)
    duplicates = conflicts = 0
    for group in groups.values():
        if len(group) > 1:
            duplicates += len(group) - 1
            if len({r['value'] for r in group}) > 1:
                conflicts += 1
                for row in group:
                    row['quality_flag'] = 'conflicting_duplicate'
    date_audit = calendar_alignment_audit(rows)
    checked = set()
    comparisons = date_audit['comparisons']
    for station_id in DATE_CHECKED_STATIONS:
        passed = []
        for variable, tolerance in [('air_temperature_mean', .5), ('precipitation_total', 2.)]:
            for season in ['Jan-Sep', 'Oct-Dec']:
                candidates = {r['year_shift']: r for r in comparisons if r['station_id'] == station_id
                              and r['variable'] == variable and r['season'] == season}
                selected = candidates.get(0)
                passed.append(bool(selected and selected['n'] >= 24 and selected['mae'] < tolerance
                    and all(selected['mae'] < r['mae'] for shift,r in candidates.items() if shift != 0)))
        if all(passed):
            checked.add(station_id)
    date_audit['checked_station_ids'] = sorted(checked)
    for row in rows:
        if row['date_status'] == 'cross_checked_ending_year' and row['entity_id'] not in checked:
            row['date_status'] = 'inferred_ending_year'
    write_csv(OBSERVATIONS_CSV, list(rows[0]) if rows else [], rows)

    suspects = [row for row in rows if row["quality_flag"] != "ok"]
    matched = sum(1 for row in rows if row["entity_id"])
    fuzzy_matched = sum(1 for row in rows if row["match_method"] == "fuzzy")

    # The JSON projection the browser reads: indexed by station and variable,
    # one array of twelve months per calendar year.
    series: dict = {}
    for row in rows:
        key = row["entity_id"] or f"unmatched:{row['source_file']}:{row['source_sheet']}:{row['station_name_latin']}"
        station = series.setdefault(key, {
            "label": row["station_name_latin"], "entityId": row["entity_id"],
            "regionEn": row["region_en"], "variables": {},
        })
        variable = station["variables"].setdefault(row["variable"], {"unit": row["unit"], "years": {}})
        months = variable["years"].setdefault(str(row["calendar_year"]), [None] * 12)
        slot = row["month"] - 1
        months[slot] = float(row["value"]) if row["quality_flag"] == "ok" else None
    write_json(OBSERVATIONS_JSON, {
        "version": "1.0", "generatedAt": retrieved,
        "note": "Monthly regional station observations, one array of twelve calendar months per year. "
                "null means missing or quarantined. Duplicate identical values are represented once. "
                "Air/precipitation dates use ending years, cross-checked for Tashkent/Pskem and inferred elsewhere.",
        "stations": series,
    })

    variables = sorted({row["variable"] for row in rows})
    years = sorted({int(row["calendar_year"]) for row in rows})
    manifest = {
        "version": "1.0",
        "generatedAt": retrieved,
        "observationClass": "station_observation",
        "spatialScope": "station",
        "sources": [{"file": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()} for path in sorted(AIR_TEMP_DIR.glob("*.xlsx"))] + [
            {"file": SOIL_TEMP_PATH.name, "sha256": hashlib.sha256(SOIL_TEMP_PATH.read_bytes()).hexdigest() if SOIL_TEMP_PATH.exists() else None},
        ],
        "counts": {
            "monthlyValues": len(rows),
            "variables": len(variables),
            "stationsMatchedToNetwork": len({row["entity_id"] for row in rows if row["entity_id"]}),
            "valuesWithMatchedStation": matched,
            "valuesWithUnmatchedStation": len(rows) - matched,
            "valuesMatchedByFuzzyName": fuzzy_matched,
            "suspectValues": len(suspects),
            "duplicateExtraRows": duplicates,
            "conflictingDuplicateKeys": conflicts,
        },
        "variables": variables,
        "calendarAlignment": date_audit,
        "temporalCoverage": {"start": years[0] if years else None, "end": years[-1] if years else None},
        "suspectBlockLabels": {
            "count": len(swaps),
            "sheets": sorted(swaps),
            "rule": "No automatic swaps. Negative precipitation quarantines the sheet pending source review.",
        },
        "unmatchedStations": sorted(unmatched),
        "qualityRanges": RANGES,
        "notes": [
            "Air/precipitation use hydrological ending years; cross-checked for Tashkent/Pskem against existing calendar observations, inferred elsewhere. Source_cell and date_status retain traceability.",
            "Soil temperature is not soil texture or soil type; measurement depth is not supplied.",
            "Fuzzy name candidates are not accepted as station identities. Network identity is not proof of ontology membership.",
            "quality_flag='suspect_out_of_range' values are kept, not deleted; filter on quality_flag=='ok' for "
            "an implementation-ready series.",
            "Pedya index columns exist in every air_temp_precip sheet but carry no data in this delivery; no "
            "ped_drought_index rows are published as a result.",
            "Station names were repaired for Cyrillic/Latin homoglyphs and matched to "
            "hydromet-meteo-network.csv by their transliterated form; entity_id is blank where no match was found.",
        ],
    }
    write_json(MANIFEST, manifest)

    print(f"\n{len(rows):,} monthly values, {len(suspects):,} flagged suspect, "
          f"{matched:,}/{len(rows):,} matched to a station entity, {len(swaps)} sheet(s) quarantined")
    if unmatched:
        print(f"{len(unmatched)} station blocks did not match the network (see manifest.unmatchedStations)")


if __name__ == "__main__":
    main()
