"""Export a cited interpretation, figure atlas and supporting workbook."""
import csv
import hashlib
import json
from datetime import datetime
from pathlib import Path
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from build_chirchik_case_studies import OUT, read_csv, write_json

SOURCES=[
 ('ERA5-Land monthly aggregates','Copernicus / ECMWF, Earth Engine catalogue','https://developers.google.com/earth-engine/datasets/catalog/ECMWF_ERA5_LAND_MONTHLY_AGGR','Air temperature, precipitation, evaporation, radiation and runoff'),
 ('MOD11A2 v6.1','NASA LP DAAC, Earth Engine catalogue','https://developers.google.com/earth-engine/datasets/catalog/MODIS_061_MOD11A2','1 km daytime surface temperature; 8-day composites'),
 ('MOD16A2GF v6.1','NASA LP DAAC, Earth Engine catalogue','https://developers.google.com/earth-engine/datasets/catalog/MODIS_061_MOD16A2GF','500 m terrestrial ET; 8-day totals and QA'),
 ('MCD43A3 v6.1','NASA LP DAAC, Earth Engine catalogue','https://developers.google.com/earth-engine/datasets/catalog/MODIS_061_MCD43A3','500 m albedo; daily rolling 16-day retrievals'),
 ('Copernicus DEM GLO-30 2024_1','Copernicus, Earth Engine catalogue','https://developers.google.com/earth-engine/datasets/catalog/COPERNICUS_DEM_GLO30_2024_1','30 m surface elevation'),
 ('Latest land-cover data release','Esri / Impact Observatory, 2026','https://www.esri.com/about/newsroom/arcnews/latest-land-cover-data-release-shows-more-change-over-time','Annual Sentinel-2 land cover through 2025'),
 ('LULC methodology and accuracy','Impact Observatory','https://www.impactobservatory.com/legal/lulc-methodology-accuracy.pdf','Classification methods and global accuracy scope'),
 ('Cloud removal methodology from MODIS snow cover product','Gafurov and Bárdossy, HESS, 2009','https://hess.copernicus.org/articles/13/1361/2009/','MODSNOW-style methodology context'),
 ('Permutation feature importance','scikit-learn documentation','https://scikit-learn.org/stable/modules/permutation_importance','Predictive reliance and correlated features'),
 ('RandomForestRegressor','scikit-learn documentation','https://scikit-learn.org/stable/modules/generated/sklearn.ensemble.RandomForestRegressor.html','Model specification'),
 ('Inherent benchmark or not? Comparing NSE and KGE','Knoben et al., HESS, 2019','https://hess.copernicus.org/articles/23/4323/2019/','Evaluation metric interpretation'),
 ('Delivered Pskem station workbooks','Local project observations; original filenames and hashes in chirchik.manifest.json','','Ground P/T/Q and snow-day comparisons'),
]
NAMES={'seasonal_climatology':'Seasonal baseline','physical_bucket':'Snow–soil bucket','random_forest':'Random forest','bayesian_linear':'Bayesian linear'}
BLUE,ORANGE,PURPLE='#177e9b','#b86b28','#8263ad'


def figure_atlas(d,a):
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9,'axes.spines.top':False,'axes.spines.right':False,'axes.titlelocation':'left','axes.titleweight':'bold'})
    paths=[]
    with PdfPages(OUT/'chirchik-scientific-atlas.pdf') as pdf:
        def save(fig,name,caption):
            fig.supxlabel(caption,fontsize=8)
            pdf.savefig(fig)
            fig.savefig(OUT/f'{name}.png',dpi=180)
            paths.append(f'{name}.png');plt.close(fig)
        meta=d['profile_metadata'];(south,west),(north,east)=meta['map_bounds']
        fig,ax=plt.subplots(figsize=(11.7,8.3),layout='constrained')
        ax.imshow(plt.imread(OUT/'pskem-elevation-overview.png'),extent=[west,east,south,north])
        ax.imshow(plt.imread(OUT/'pskem-terrain-overview.png'),extent=[west,east,south,north],alpha=.20)
        geo=json.loads((OUT/'pskem-candidate-catchment.geojson').read_text())
        for feature in geo['features']:
            geometry=feature['geometry'];polygons=geometry['coordinates'] if geometry['type']=='MultiPolygon' else [geometry['coordinates']]
            for polygon in polygons:
                xy=np.array(polygon[0]);ax.plot(xy[:,0],xy[:,1],color='#274f5d',lw=.6)
        base=json.loads((OUT/'chirchik.json').read_text());stations={r['station']:r for r in base['inventory']}
        for name,r in stations.items():
            if west<float(r['longitude'])<east and south<float(r['latitude'])<north:
                ax.scatter(float(r['longitude']),float(r['latitude']),c=ORANGE,edgecolor='white',s=55,zorder=4);ax.annotate(name,(float(r['longitude']),float(r['latitude'])),xytext=(7,7),textcoords='offset points',fontsize=9)
        c=base['catchment'];ax.scatter(c['gauge_longitude'],c['gauge_latitude'],marker='s',s=70,c='#a42b47',edgecolor='white',zorder=4);ax.annotate('Mullala gauge\nlocation unresolved',(c['gauge_longitude'],c['gauge_latitude']),xytext=(12,-20),textcoords='offset points',fontsize=9)
        ax.annotate('N',xy=(.94,.93),xytext=(.94,.83),xycoords='axes fraction',ha='center',arrowprops={'arrowstyle':'->','lw':1.5},fontsize=12)
        from pyproj import Geod
        lon0,lat0=west+.08,south+.035;lon1,lat1,_=Geod(ellps='WGS84').fwd(lon0,lat0,90,10000)
        ax.plot([lon0,lon1],[lat0,lat1],color='#142e3b',lw=3);ax.text((lon0+lon1)/2,lat0+.012,'10 km',ha='center',fontsize=9)
        from matplotlib.colors import LinearSegmentedColormap,Normalize
        ramp=LinearSegmentedColormap.from_list('elevation',['#416a53','#a0aa72','#d3b285','#a28778','#f2f3f2'])
        fig.colorbar(plt.cm.ScalarMappable(norm=Normalize(800,4500),cmap=ramp),ax=ax,shrink=.65,label='Copernicus elevation (m); terrain shading overlaid')
        ax.set(xlim=(west,east),ylim=(south,north),xlabel='Longitude (°E)',ylabel='Latitude (°N)',title='Pskem mountain study domain and observation locations');ax.set_aspect(1/np.cos(np.radians((south+north)/2)))
        save(fig,'pskem-study-map','Sources: Copernicus DEM GLO-30 2024_1; supplied station coordinates; HydroBASINS level-12 trace.\nThe boundary includes the complete outlet unit and is provisional. Tashkent lies outside this mountain map.')
        profiles=[{k:float(v) for k,v in r.items()} for r in d['elevation_profiles']]
        z=[r['minimum_m']+125 for r in profiles]
        fig,axes=plt.subplots(2,2,figsize=(11.7,8.3),layout='constrained')
        for ax,key,title,unit in zip(axes.flat,['area_km2','temperature_c','precipitation_mm_year','cumulative_runoff_above_mcm_year'],['a. Mountain area','b. Mean air temperature','c. Annual precipitation','d. Cumulative runoff generated above'],['km²','°C','mm/year','million m³/year']):
            ax.plot(z,[r[key] for r in profiles],marker='o',ms=3,color=BLUE);ax.set(title=title,xlabel='Elevation-band midpoint (m)',ylabel=unit);ax.grid(alpha=.2)
        fig.suptitle('Pskem: terrain and climate by 250 m elevation band',fontsize=15)
        save(fig,'pskem-elevation-profiles','Sources: Copernicus DEM 30 m; ERA5-Land 2010–2024. Coarse climate sampled across terrain bands; not fine-scale downscaling.\nCumulative runoff is modelled generation, not routed or measured discharge. Provisional catchment.')
        fig,axes=plt.subplots(1,2,figsize=(11.7,8.3),layout='constrained');classes=d['profile_metadata']['classes'];colors=d['profile_metadata']['class_colors']
        for ax,year in zip(axes,[2017,2025]):
            land=[r for r in d['landcover_elevation'] if int(r['year'])==year and r['valid_class']=='True'];left=np.zeros(len(z))
            total=np.array([sum(float(r['area_km2']) for r in land if float(r['minimum_m'])==p['minimum_m']) for p in profiles])
            for code,label in classes.items():
                if code=='10':continue
                values=np.array([sum(float(r['area_km2']) for r in land if str(r['class_code'])==code and float(r['minimum_m'])==p['minimum_m']) for p in profiles])
                percent=np.divide(values,total,out=np.zeros_like(values),where=total>0)*100
                ax.barh(z,percent,left=left,height=205,color=colors[code],label=label);left+=percent
            ax.set(title=str(year),xlabel='Share of valid classified area (%)',ylabel='Elevation-band midpoint (m)',xlim=(0,100))
        axes[1].legend(loc='lower right',fontsize=8,framealpha=.95)
        fig.suptitle('Pskem: 10 m land cover along the mountain',fontsize=15)
        save(fig,'pskem-landcover-elevation','Source: Esri / Impact Observatory / Microsoft, community-hosted Earth Engine series.\nChanges combine real change and classification differences. No local independent accuracy sample is available.')
        fig,axes=plt.subplots(2,1,figsize=(11.7,8.3),layout='constrained')
        rows=[r for r in a['snow_monthly'] if r['sensor']=='combined' and r['elevation_band']=='all'];years=sorted({int(r['period'][:4]) for r in rows});lookup={r['period']:r for r in rows}
        for ax,key,title,cmap in zip(axes,['snow_cover_percent','mean_valid_area_percent'],['a. Snow among usable pixels','b. Mean daily usable area'],['Blues','YlGn']):
            matrix=np.array([[lookup.get(f'{year}-{m:02}',{}).get(key) for m in range(1,13)] for year in years],dtype=float)
            color=plt.get_cmap(cmap).copy();color.set_bad('#b8bec4');im=ax.imshow(matrix,aspect='auto',vmin=0,vmax=100,cmap=color)
            ax.set(title=title,xticks=range(12),xticklabels=['J','F','M','A','M','J','J','A','S','O','N','D'],yticks=range(0,len(years),2),yticklabels=years[::2]);fig.colorbar(im,ax=ax,label='%')
        fig.suptitle('MODSNOW-style monitoring: snow and observation completeness',fontsize=15)
        save(fig,'pskem-snow-calendar','Sources: MOD10A1 + MYD10A1 v6.1. Same-day Terra priority, Aqua fallback; no temporal interpolation.\nSnow means need ≥70% daily valid area and ≥10 qualifying days/month. Grey means missing or insufficient observations.')
        rows=[r for r in d['model_series'] if r['phase']=='held_out'];dates=[datetime.strptime(r['period'],'%Y-%m') for r in rows]
        fig,axes=plt.subplots(3,1,figsize=(11.7,9.5),layout='constrained')
        for ax,key in zip(axes,['physical_bucket','random_forest','bayesian_linear']):
            ax.plot(dates,[r['observed_cms'] for r in rows],color=BLUE,label='Observed',lw=1.4)
            ax.plot(dates,[r[key] for r in rows],color=ORANGE,label=NAMES[key],lw=1.2)
            if key=='bayesian_linear':ax.fill_between(dates,[r['bayesian_lower95'] for r in rows],[r['bayesian_upper95'] for r in rows],color=PURPLE,alpha=.18,label='95% predictive interval')
            score=next(m['scores'] for m in d['monthly_models'] if m['model']==key)
            ax.set(title=f"{NAMES[key]} · RMSE {score['rmse']:.1f} m³/s · NSE {score['nse']:.3f}",ylabel='m³/s');ax.legend(ncol=3,frameon=False,fontsize=8);ax.grid(alpha=.2)
        fig.suptitle('Monthly model verification on unseen 2011–2017 observations',fontsize=15)
        save(fig,'pskem-model-comparison','Sources: delivered screened discharge; ERA5-Land forcing. Training: 2001–2010, warm-up: 2000.\nBayesian intervals omit boundary and serial-error uncertainty. Experimental hindcasts for a provisional catchment.')
        fig,axes=plt.subplots(1,2,figsize=(11.7,6.5),layout='constrained')
        models=d['monthly_models'];axes[0].barh([NAMES[r['model']] for r in models],[r['scores']['rmse'] for r in models],color=[BLUE,ORANGE,PURPLE,'#618c6b']);axes[0].invert_yaxis();axes[0].set(title='a. Typical test error: lower is better',xlabel='RMSE (m³/s)')
        imp=d['feature_importance'];axes[1].barh([r['variable'].replace('_',' ') for r in imp],[r['mean_rmse_increase_cms'] for r in imp],color=ORANGE);axes[1].axvline(0,color='#555',lw=.7);axes[1].invert_yaxis();axes[1].set(title='b. Forest reliance on interannual inputs',xlabel='Mean change in test RMSE (m³/s)')
        fig.suptitle('Prediction skill and variable importance',fontsize=15)
        save(fig,'pskem-model-skill','Permutation rearranges whole held-out years, preserving month alignment (20 repeats). Positive bars mean larger prediction errors.\nSeasonal inputs repeat every year; zero here cannot establish that calendar season is unimportant. Association is not causation.')
        surface={p:{r['period']:r for r in d['surface_monthly'] if r['product']==p} for p in ['lst','et','albedo']};energy=d['energy_monthly']
        dates=[datetime.strptime(r['period'],'%Y-%m') for r in energy];air=[float(r['temperature_c']) for r in energy];lst=[surface['lst'].get(r['period'],{}).get('value',np.nan) for r in energy]
        fig,axes=plt.subplots(2,2,figsize=(11.7,8.3),layout='constrained')
        axes[0,0].plot(dates,air,lw=.8,label='ERA5 air',color=BLUE);axes[0,0].plot(dates,lst,lw=.8,label='MODIS daytime surface',color=ORANGE);axes[0,0].set(title='a. Two different temperatures',ylabel='°C');axes[0,0].legend(fontsize=8)
        points=np.array([(x,y) for x,y in zip(air,lst) if y is not None and np.isfinite(y)]);slope,intercept=np.polyfit(points[:,0],points[:,1],1);xx=np.linspace(points[:,0].min(),points[:,0].max(),100)
        axes[0,1].scatter(points[:,0],points[:,1],s=8,color=BLUE,alpha=.5);axes[0,1].plot(xx,slope*xx+intercept,color=ORANGE);axes[0,1].set(title=f'b. Descriptive fit: surface = {slope:.2f} × air + {intercept:.2f}',xlabel='ERA5 monthly air (°C)',ylabel='MODIS daytime surface (°C)')
        axes[1,0].plot(dates,[float(r['solar_down_wm2']) for r in energy],lw=.8,color=ORANGE,label='Incoming solar');axes[1,0].plot(dates,[float(r['net_solar_wm2']) for r in energy],lw=.8,color=PURPLE,label='Net solar');axes[1,0].set(title='c. Solar energy',ylabel='W/m²');axes[1,0].legend(fontsize=8)
        axes[1,1].plot(dates,[surface['albedo'].get(r['period'],{}).get('value',np.nan) for r in energy],lw=.8,color=PURPLE);axes[1,1].set(title='d. Black-sky shortwave albedo',ylabel='Reflected fraction',ylim=(0,1))
        fig.suptitle('Pskem surface energy: seasonal relationships and different footprints',fontsize=15)
        save(fig,'pskem-surface-energy','Sources: ERA5-Land, MOD11A2 and MCD43A3 v6.1. Clear-sky surface temperature is not mean air temperature.\nSeasonal linear association is descriptive; a high correlation is not independent validation.')
        fig,axes=plt.subplots(2,1,figsize=(11.7,7.5),layout='constrained')
        axes[0].plot(dates,[float(r['evapotranspiration_mm']) for r in energy],color=BLUE,lw=.9,label='ERA5 basin ET');axes[0].plot(dates,[surface['et'].get(r['period'],{}).get('value',np.nan) for r in energy],color=ORANGE,lw=.9,label='MOD16 valid terrestrial footprint');axes[0].set(title='a. Monthly evapotranspiration',ylabel='mm/month');axes[0].legend()
        axes[1].plot(dates,[surface['et'].get(r['period'],{}).get('mean_valid_area_percent',np.nan) for r in energy],color=ORANGE,lw=.9);axes[1].set(title='b. Area represented by the MOD16 mean',ylabel='% of candidate catchment',ylim=(0,100))
        fig.suptitle('Evapotranspiration: a comparison needs matching spatial support',fontsize=15)
        save(fig,'pskem-evapotranspiration','Sources: ERA5-Land and MOD16A2GF v6.1. MOD16 excludes much water, barren and snow/ice area.\nEight-day ET totals are allocated by overlapping calendar days; incomplete months remain gaps. These are model estimates, not independent ET measurements.')
        fig,axes=plt.subplots(1,2,figsize=(11.7,6.5),layout='constrained')
        corrections=[r for r in a['climate_validation'] if r['correction'] and r['station']=='Pskem']
        for ax,var,unit in zip(axes,['air_temperature_mean','precipitation_total'],['°C','mm/month']):
            records=[r for r in corrections if r['variable']==var];xx=np.arange(len(records))
            ax.bar(xx-.18,[r['correction']['raw_test']['rmse'] for r in records],width=.36,label='Raw',color=ORANGE);ax.bar(xx+.18,[r['correction']['corrected_test']['rmse'] for r in records],width=.36,label='Training-only correction',color=BLUE);ax.set(xticks=xx,xticklabels=[r['product'] for r in records],ylabel=f'RMSE ({unit})',title='Temperature' if var.startswith('air') else 'Precipitation');ax.legend(fontsize=8)
        fig.suptitle('Pskem forcing correction tested on 2018–2024',fontsize=15)
        save(fig,'pskem-climate-validation','Sources: station observations, ERA5-Land, CHIRPS v3. Monthly corrections fitted only through 2017.\nA successful empirical correction does not explain the original bias or validate transfer to other elevations.')
        fig,axes=plt.subplots(1,2,figsize=(11.7,6.5),layout='constrained');water=a['reservoir'];r=water['monthly'];dates=[datetime.strptime(v['period'],'%Y-%m') for v in r]
        axes[0].plot(dates,[v['water_area_km2'] if v['eligible'] else np.nan for v in r],color=BLUE,lw=1);axes[0].set(title='a. Screened JRC water detections',ylabel='km²')
        pairs=water['sentinel_check']['pairs'];axes[1].scatter([p['observed'] for p in pairs],[p['predicted'] for p in pairs],color=BLUE);axes[1].plot([10,40],[10,40],color=ORANGE,ls='--');axes[1].set(title=f"b. {len(pairs)} same-month sensor pairs",xlabel='JRC / Landsat (km²)',ylabel='Sentinel-2 (km²)')
        fig.suptitle('Charvak: water extent and an independent-sensor check',fontsize=15)
        save(fig,'charvak-water-verification','Sources: JRC Global Surface Water 1.4 and Sentinel-2 SR Harmonized. Both comparisons require ≥95% valid area.\nAcquisition dates differ within the month. The polygon plus 1 km domain includes adjacent channels; area is not storage volume.')
    return paths


def report(d,a,paths):
    m={r['model']:r['scores'] for r in d['monthly_models']};improvement=100*(1-m['random_forest']['rmse']/m['seasonal_climatology']['rmse'])
    p=[{k:float(v) for k,v in r.items()} for r in d['elevation_profiles']];area=sum(r['area_km2'] for r in p);peak=max(p,key=lambda r:r['area_km2'])
    fit=d['physical_parameters'];bayes=d['bayesian_interval_check'];fresh=d['freshness'];snow=a['seasonal_flow'];water=a['reservoir']['sentinel_check']
    lines=['# Pskem mountain water: terrain, energy and model verification','',
        f'The random forest produces the lowest monthly discharge RMSE among four tested specifications: {m["random_forest"]["rmse"]:.2f} m³/s, a {improvement:.1f}% reduction from the seasonal baseline on 84 held-out months. The physical snow–soil bucket conserves water but has higher prediction error; the Bayesian model supplies explicit uncertainty with a wide predictive range. These results support an exploratory model comparison, not an operational release.','',
        'The principal spatial problem is unresolved gauge identity. The stored coordinate is almost on a small tributary with 11.7 km² reference upstream area, whereas plausible Pskem main-stem reaches lie about 700 m away. The provisional domain contains the entire outlet level-12 unit. Every basin profile and calibrated model must be recomputed if the verified outlet changes.','',
        '## Terrain and land-cover evidence','',
        f'The Copernicus elevation-band reduction represents {area:.1f} km². The largest 250 m band is {peak["minimum_m"]:.0f}–{peak["maximum_m"]:.0f} m, containing {peak["area_km2"]:.1f} km². Charts show area, mean air temperature, annual precipitation, local runoff generation, and runoff accumulated from higher bands. The DEM is a 30 m digital surface model. It updates the terrain source but does not provide finer nominal resolution than SRTM 30 m.[^5]','',
        'The 2010–2024 climate summary intersects coarse ERA5 cells with fine terrain bands. Temperature is the mean of 180 monthly means; annual precipitation and runoff use fifteen annual-equivalent totals. These profiles describe spatial association in the reanalysis. They are not station-derived lapse rates and do not contain independent 250 m climate information. Runoff depth multiplied by band area gives local generated volume; summing from high elevations downward does not route rivers or measure discharge at each elevation.[^1]','',
        'Esri/Impact Observatory classifications are available for 2017–2025 and were reduced at 10 m by elevation band. Rangeland, bare ground and seasonal snow/ice are especially relevant mountain classes. The stacked charts show class shares among valid mapped pixels. Cloud/unclassified areas are excluded and retained separately in the CSV. A mapped snow/ice class is not a glacier inventory, and class switches between years do not alone establish ecological change.[^6][^7]','',
        'No verified, openly licensed sub-30 m DEM was identified for this catchment. Resampling a 30 m model to 10 m would create smaller pixels without new terrain observations. A professional upgrade should compare surveyed control points or independently acquired stereo/lidar terrain before changing flow paths or claiming finer accuracy. Inter-product elevation spread is not error against ground truth.','',
        '## Temperature, radiation, albedo and evapotranspiration','',
        'The surface-energy charts separate two temperatures: ERA5 monthly 2 m air temperature and MODIS daytime clear-sky surface skin temperature. The latter measures the radiometric surface, which can heat much faster than the air. A fitted surface-versus-air line is descriptive and strongly affected by season, time of day and clear-sky sampling. It cannot validate station air temperature.[^1][^2]','',
        'Incoming solar radiation supplies energy; net solar radiation subtracts reflected shortwave energy. Monthly accumulated joules per square metre are divided by seconds in the calendar month to obtain mean watts per square metre. Albedo is the reflected fraction. MODIS black-sky shortwave albedo is a modelled directional illumination quantity, not a complete measured surface-energy balance. Daily MCD43 estimates use overlapping 16-day retrieval windows.[^1][^4]','',
        'Evapotranspiration combines water evaporating from surfaces and transpiring through vegetation. ERA5 ET is a basin reanalysis estimate. MOD16 ET is a terrestrial model with quality-filtered coverage that excludes much barren, water and snow/ice area. Its monthly means therefore represent a changing subset of this mountain basin. Eight-day ET totals are allocated to calendar months in proportion to overlapping days; months without all composite days remain blank. LST composite means receive duration weighting; albedo uses available daily retrievals. Similar-looking curves do not establish agreement on a common footprint.[^3]','',
        '## Physical model','',
        'Each monthly step divides precipitation between rain and snow with a smooth temperature-dependent fraction. Snow accumulates in a store. Positive monthly temperature produces potential degree-day melt, capped by snow availability. Rain plus melt feeds quick flow and a soil store; evapotranspiration is limited by soil water, excess spills to runoff, and a fraction of remaining soil water drains as baseflow. The model is an explicitly simplified numerical bucket, not a full energy-balance or glacier model.','',
        'The water balance is P − ET − Q − Δ(snow + soil) = 0 in equivalent depth. Runoff is converted to m³/s with the provisional catchment area and actual calendar-month length. Monthly mean temperature cannot reconstruct daily freezing crossings; melt timing is correspondingly approximate. Potential ET is clamped at zero when its signed conversion is negative. There is no explicit channel travel time, reservoir rule or glacier ice component.','',
        '| Parameter | Fitted value | Search bounds |','| --- | ---: | --- |']
    for key,value in fit.items():lines.append(f'| {key} | {value:.4f} | {d["physical_parameter_bounds"][key]} |')
    lines+=['',f'Calibration uses differential evolution with a fixed seed, 60 maximum iterations and 2000 warm-up. The objective sees only 2001–2010 observations. The soil capacity reaches its 500 mm upper bound, indicating that the fit is constrained by the search space and may compensate for omitted processes. Maximum absolute water-balance residual is {d["physical_mass_balance_max_abs_mm"]:.2e} mm. Numerical conservation is a check on implementation, not proof that each water source is correctly identified.','',
        '## Machine learning and Bayesian alternatives','',
        'The random forest uses 200 trees, maximum depth five and at least eight training observations per leaf. Predictors are current precipitation, air temperature, potential ET, preceding-three-month precipitation and sine/cosine calendar season. These hyperparameters were fixed before scoring held-out years. Current-month inputs make this a retrospective reconstruction, not a forecast issued before the month starts. Nonlinear methods cannot be assumed to extrapolate to unprecedented climate conditions.[^10]','',
        'Variable importance rearranges whole held-out years of one input while retaining calendar-month alignment, then reports the change in RMSE over twenty repeats. Antecedent precipitation has the largest positive mean change in this experiment. Calendar predictors are identical across years, so this test cannot measure their importance. Very small or negative changes for the other covariates do not show that their physical processes are unimportant. Correlated predictors can substitute for each other; these are predictive associations, not causal effects.[^9]','',
        'The Bayesian alternative is conjugate linear regression, which can be solved analytically. Predictors and response are standardized using the training period alone. Coefficients conditional on residual variance have a zero-mean normal prior with precision one; the intercept precision is 10⁻⁶. Residual variance has an inverse-gamma prior with shape two and scale one. Updating these assumptions with training observations gives a Student-t posterior predictive distribution. No MCMC chains are required for this conjugate model.','',
        f'The nominal 95% interval contains {100*bayes["held_out_coverage"]:.1f}% of the held-out observations, with mean width {bayes["mean_width_cms"]:.1f} m³/s. The large width matters as much as coverage. Gaussian predictions and lower intervals can be negative and are retained visibly; a later positive-response model would require a new predeclared evaluation. Serial error dependence, rating-curve uncertainty, climate-product uncertainty and gauge-boundary uncertainty are omitted.','',
        '## Equal-period verification','',
        '| Model | Test months | RMSE (m³/s) | Bias (m³/s) | NSE | KGE |','| --- | ---: | ---: | ---: | ---: | ---: |']
    for key,s in m.items():lines.append(f'| {NAMES[key]} | {s["n"]} | {s["rmse"]:.2f} | {s["bias"]:.2f} | {s["nse"]:.3f} | {s["kge"]:.3f} |')
    lines+=['','RMSE measures typical squared-error magnitude; smaller is better. Bias is prediction minus observation. NSE equal to one is perfect and zero matches the observed test-period mean. KGE combines correlation, variability and mean bias, and has a different benchmark interpretation. The same screened months are used for every monthly model. Seven test years do not justify strong claims about generalisation to future climates; the ranking is descriptive and no claim of statistically significant superiority is made.[^11]','',
        f'The separate April 1 experiment uses {len(snow["training_years"])} eligible training years and {len(snow["test_years"])} eligible test years under the strict March snow-coverage gate. Adding the chosen snow-cover predictor worsens seasonal-volume RMSE relative to climate-only predictors on that identical cohort. This negative result is retained. It does not establish that snow is hydrologically unimportant: coverage, predictor timing, a tiny sample and model specification limit the experiment. See the companion historical-validation report for every specification and exclusion.','',
        '## MODSNOW-style interpretation','',
        'The implementation combines same-day quality-screened Terra and Aqua snow observations, retaining Terra where valid and using Aqua otherwise. Daily snow area is stratified by elevation, and monthly calendars show both snow fraction and usable observation coverage. NDSI is classified explicitly; it is not treated as fractional snow cover. Daily reductions require at least 70% valid area and monthly snow summaries require ten qualifying days. Station snow-day counts are compared with minimum/maximum bounds caused by unknown days.','',
        'This is MODSNOW-style monitoring, not a reproduction of the full published MODSNOW cloud-removal algorithm. No neighbouring dates or inferred elevation rules fill the cloud gaps. A professional cloud-removal extension should hide known clear pixels, reconstruct them with only information available at the intended issue time, and score the withheld labels by elevation and season. That experiment is needed before filling unknown snow states and using them as if observed.[^8]','',
        f'Charvak supplies another independent-sensor check: {water["scores"]["n"]} screened month pairs between JRC/Landsat and Sentinel-2 yield RMSE {water["scores"]["rmse"]:.2f} km². Different acquisition dates and shared optical limitations remain. This validates neither absolute storage nor an area–volume relation. The supplied ground records do not contain a stage/discharge rating curve or reservoir bathymetry, so no physical rating curve is invented.','',
        '## Freshness and the readiness bar','',
        f'ERA5 forcing and experimental monthly continuation reach {fresh["era5_latest_available"]}. The latest source dates are MODIS LST {fresh["modis_latest_available"]["lst"]}, MOD16 ET {fresh["modis_latest_available"]["et"]}, and albedo {fresh["modis_latest_available"]["albedo"]}. Combined snow reaches {max(r["period"] for r in a["snow_monthly"] if r["sensor"]=="combined")}; a partial latest month is not a complete calendar month. Historical station observations retain their own end dates.','',
        'The coloured bar is an evidence checklist. Green means the named check is complete; amber means an explicit limitation remains; red means a required problem is unresolved. Overall release becomes green only when every listed check is green. It never averages a red spatial problem into a reassuring percentage. The current result is red. Updating the remote inputs is already possible; 2018 onward is labelled unverified continuation because no corresponding discharge observations are supplied.','',
        '| Check | State | Evidence or remaining issue |','| --- | --- | --- |']
    for r in d['readiness']['steps']:lines.append(f'| {r["label"]} | {r["status"]} | {r["reason"]} |')
    lines+=['','## Further professional analyses, in order of value','',
        'First, resolve gauge/reach identity and obtain recent discharge with rating-curve metadata. Then repeat the extractions for the reviewed partial-outlet catchment and plausible neighbouring boundaries. Boundary sensitivity, measurement uncertainty and a genuine recent temporal holdout are more informative than simply adding a more complex learner.','',
        'Second, obtain daily temperature and precipitation for an elevation-band snow/rain and degree-day model. Compare its simulated snow with QA-screened satellite area using a snow depletion relation, keeping calibration and independent snow checks separate. Daily radiation can support an energy-balance melt experiment once albedo, humidity, wind and cloud assumptions are explicit. Monthly radiation and temperature alone do not constrain a detailed melt-energy balance.','',
        'Third, run nested rolling-origin model evaluation and year-block uncertainty for error differences. A hierarchical Bayesian extension can pool station/elevation biases while retaining station-specific effects. A positive discharge likelihood, autocorrelated residuals, posterior predictive checks and prior sensitivity would address clear limitations of the present linear model. The available monthly sample supports a parsimonious model more readily than a large deep-learning network.','',
        'Fourth, establish land-cover change accuracy with stratified independent labelled samples for stable and changing classes, then estimate error-adjusted change areas. Vegetation phenology, irrigation and drought indices could be compared within these verified classes. Increasing spatial resolution does not compensate for unmeasured classification errors.','',
        'Finally, glacier attribution requires dated outlines and independent mass-balance constraints; reservoir storage requires independent levels and bathymetry or a supported elevation–area–volume relation. Radar can complement cloudy optical shoreline observations after local shadow, layover, wind and ice checks. These additions address identifiable data gaps rather than treating agreement between two models as ground truth.','',
        '## Figures and supporting data','',
        '[Scientific figure atlas (PDF)](chirchik-scientific-atlas.pdf) · [Supporting workbook](chirchik-analysis.xlsx) · [Historical validation report](chirchik-deep-study.md)','']
    for name in paths:lines.append(f'![{name.replace("-"," ").removesuffix(".png")}]({name})\n')
    lines+=['## Sources','']
    for i,(title,publisher,url,role) in enumerate(SOURCES,1):lines.append(f'[^{i}]: {publisher}. '+(f'[{title}]({url})' if url else title)+f'. Used for: {role}.')
    (OUT/'chirchik-environment-study.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')


def workbook(d):
    wb=Workbook();guide=wb.active;guide.title='Read me';guide.append(['Pskem–Charvak analysis','Interpretation'])
    for row in [('Status','Exploratory; gauge assignment unresolved'),('Training / test','Monthly models: 2001–2010 / 2011–2017'),('Units','Retained in column names; snow and coverage are percent'),('Sources','See Sources sheet and adjacent manifests for hashes'),('Missing values','Blank means absent or ineligible, never zero'),('Chart data','Sheets correspond to exported figures; no daily snow rows are hidden behind monthly means')]:guide.append(row)
    files=['monthly-model-predictions.csv','elevation-profiles.csv','landcover-elevation.csv','snow-monthly.csv','surface-monthly.csv','pskem-energy-monthly.csv','corrected-station-pairs.csv','charvak-cross-sensor-pairs.csv','charvak-jrc-monthly.csv','seasonal-flow-predictions.csv','snow-station-bounds.csv']
    for name in files:
        sheet=wb.create_sheet(name.removesuffix('.csv')[:31]);rows=read_csv(OUT/name)
        if rows:
            sheet.append(list(rows[0]))
            for row in rows:
                values=[]
                for key,value in row.items():
                    if value=='':values.append(None);continue
                    try:values.append(float(value) if key not in ['period','date','station_id','source_image','year'] else value)
                    except ValueError:values.append(value)
                sheet.append(values)
    sheet=wb.create_sheet('Model scores');sheet.append(['Model','N','RMSE m3/s','Bias m3/s','NSE','KGE'])
    for r in d['monthly_models']:sheet.append([r['model']]+[r['scores'][k] for k in ['n','rmse','bias','nse','kge']])
    sheet=wb.create_sheet('Feature importance');sheet.append(['Variable','Mean RMSE increase m3/s','Minimum','Maximum'])
    for r in d['feature_importance']:sheet.append([r[k] for k in ['variable','mean_rmse_increase_cms','min','max']])
    sheet=wb.create_sheet('Sources');sheet.append(['Title','Publisher','URL / access','Use'])
    for r in SOURCES:sheet.append(r)
    for sheet in wb:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17495B')
        for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=min(65,max(16,len(str(col[0].value))+3))
    wb.save(OUT/'chirchik-analysis.xlsx')


def main():
    d=json.loads((OUT/'environment-modelling.json').read_text(encoding='utf-8'));a=json.loads((OUT/'advanced-validation.json').read_text(encoding='utf-8'))
    paths=figure_atlas(d,a);report(d,a,paths);workbook(d)
    write_json(OUT/'scientific-artifacts.manifest.json',{'inputs':[{'file':name,'sha256':hashlib.sha256((OUT/name).read_bytes()).hexdigest()} for name in ['environment-modelling.json','advanced-validation.json']],
        'script_sha256':hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),'outputs':paths+['chirchik-scientific-atlas.pdf','chirchik-environment-study.md','chirchik-analysis.xlsx']})
    print('Exported ten scientific figures, PDF atlas, cited report and chart-data workbook')


if __name__=='__main__':main()
